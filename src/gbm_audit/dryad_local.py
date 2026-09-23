"""Local-only acquisition, schema audit and normalization for the frozen Dryad confirmatory arm.

This module deliberately separates source inspection from confirmatory evaluation.  It never
computes tracking-method performance during inventory/normalization.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
from pathlib import Path, PurePosixPath
import platform
import sys
from typing import Any
import zipfile


REQUIRED_BUNDLE = "To Generate Figures.zip"
REQUIRED_BUNDLE_MD5 = "2b70accfbb4d81d41dfb10fcefa60cbf"
REQUIRED_README = "README_for_To Generate Figures.docx"
REQUIRED_README_MD5 = "fcbd2b285b860eb18d7ce600becda06d"
CANONICAL_COLUMNS = ["experiment_id", "cell_type", "track_id", "frame", "time_min", "x_um", "y_um"]
TEXT_SUFFIXES = {".csv", ".tsv", ".txt", ".dat"}
TABLE_SUFFIXES = TEXT_SUFFIXES | {".xlsx", ".xls", ".mat"}
MAX_MEMBERS = 20000
MAX_MEMBER_UNCOMPRESSED = 2_000_000_000
MAX_TOTAL_UNCOMPRESSED = 8_000_000_000
MAX_PREVIEW_BYTES = 4_000_000
MAX_COMPRESSION_RATIO = 500.0


def _digest(path: Path, algorithm: str) -> str:
    h = hashlib.new(algorithm)
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _bytes_digest(data: bytes, algorithm: str = "sha256") -> str:
    h = hashlib.new(algorithm)
    h.update(data)
    return h.hexdigest()


def _load_json(path: Path) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"expected JSON object: {path}")
    return payload


def _manifest_expected(manifest: dict, filename: str) -> dict:
    for row in manifest.get("files", []):
        if row.get("path") == filename:
            return row
    raise ValueError(f"source manifest does not contain {filename!r}")


def _verify_local_file(path: Path, expected: dict) -> dict:
    if not path.is_file():
        raise ValueError(f"required source file missing: {path}")
    algorithm = str(expected.get("digest_type", "")).lower()
    digest = str(expected.get("digest", "")).lower()
    if algorithm not in {"md5", "sha256"} or not digest:
        raise ValueError(f"invalid frozen digest metadata for {path.name}")
    observed = _digest(path, algorithm)
    if observed.lower() != digest:
        raise ValueError(
            f"source digest mismatch for {path.name}: expected {algorithm}:{digest}, observed {observed}"
        )
    expected_size = expected.get("size_bytes")
    if isinstance(expected_size, int) and path.stat().st_size != expected_size:
        raise ValueError(
            f"source size mismatch for {path.name}: expected {expected_size}, observed {path.stat().st_size}"
        )
    return {
        "filename": path.name,
        "size_bytes": path.stat().st_size,
        "verified_digest_type": algorithm,
        "verified_digest": observed,
        "sha256": _digest(path, "sha256"),
    }


def verify_source_dir(source_dir: Path, source_manifest_path: Path, *, require_readme: bool = True) -> dict:
    manifest = _load_json(source_manifest_path)
    bundle_expected = _manifest_expected(manifest, REQUIRED_BUNDLE)
    if bundle_expected.get("digest") != REQUIRED_BUNDLE_MD5:
        raise ValueError("repository source manifest bundle identity drifted")
    files = [_verify_local_file(source_dir / REQUIRED_BUNDLE, bundle_expected)]
    readme_path = source_dir / REQUIRED_README
    if require_readme or readme_path.exists():
        readme_expected = _manifest_expected(manifest, REQUIRED_README)
        if readme_expected.get("digest") != REQUIRED_README_MD5:
            raise ValueError("repository source manifest README identity drifted")
        files.append(_verify_local_file(readme_path, readme_expected))
    return {
        "source_manifest_sha256": _digest(source_manifest_path, "sha256"),
        "verified_files": files,
        "bundle_path": str(source_dir / REQUIRED_BUNDLE),
    }


def _validate_member(info: zipfile.ZipInfo) -> None:
    path = PurePosixPath(info.filename)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError(f"unsafe archive member path: {info.filename}")
    if info.file_size > MAX_MEMBER_UNCOMPRESSED:
        raise ValueError(f"archive member exceeds size limit: {info.filename}")
    if info.compress_size > 0 and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO:
        raise ValueError(f"archive member exceeds compression-ratio limit: {info.filename}")


def _text_preview(data: bytes, suffix: str) -> dict:
    text = data.decode("utf-8-sig", errors="replace")
    lines = [line for line in text.splitlines() if line.strip()][:8]
    if not lines:
        return {"kind": "text", "lines": []}
    if suffix == ".tsv":
        delimiter = "\t"
    elif suffix == ".csv":
        delimiter = ","
    else:
        try:
            delimiter = csv.Sniffer().sniff("\n".join(lines[:5]), delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = None
    header = []
    if delimiter:
        header = next(csv.reader([lines[0]], delimiter=delimiter), [])
    return {
        "kind": "text",
        "detected_delimiter": delimiter,
        "header": [item.strip() for item in header],
        "sample_lines": lines[:5],
    }


def _xlsx_preview(data: bytes) -> dict:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return {"kind": "xlsx", "inspection": "requires optional dependency openpyxl"}
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets[:20]:
        rows = []
        for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
            rows.append([None if value is None else str(value) for value in row[:30]])
        sheets.append({"name": sheet.title, "sample_rows": rows})
    return {"kind": "xlsx", "sheets": sheets}


def _mat_preview(data: bytes) -> dict:
    try:
        import scipy.io  # type: ignore
    except ImportError:
        return {"kind": "mat", "inspection": "requires optional dependency scipy"}
    try:
        variables = scipy.io.whosmat(io.BytesIO(data))
    except Exception as exc:  # scipy raises several format-specific exception types
        return {"kind": "mat", "inspection_error": f"{type(exc).__name__}: {exc}"}
    return {
        "kind": "mat",
        "variables": [{"name": name, "shape": list(shape), "class": klass} for name, shape, klass in variables],
    }


def inspect_bundle(bundle_path: Path) -> dict:
    members = []
    candidate_tables = []
    with zipfile.ZipFile(bundle_path) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_MEMBERS:
            raise ValueError(f"archive has too many members: {len(infos)}")
        total = 0
        for info in infos:
            _validate_member(info)
            total += info.file_size
            if total > MAX_TOTAL_UNCOMPRESSED:
                raise ValueError("archive exceeds total uncompressed-size limit")
            suffix = Path(info.filename).suffix.lower()
            row = {
                "path": info.filename,
                "suffix": suffix,
                "uncompressed_bytes": info.file_size,
                "compressed_bytes": info.compress_size,
                "crc32": f"{info.CRC:08x}",
                "candidate_tracking_table": suffix in TABLE_SUFFIXES,
            }
            members.append(row)
            if suffix in TABLE_SUFFIXES:
                preview: dict[str, Any] = {"path": info.filename, "suffix": suffix}
                if 0 < info.file_size <= MAX_PREVIEW_BYTES:
                    data = archive.read(info)
                    if suffix in TEXT_SUFFIXES:
                        preview.update(_text_preview(data, suffix))
                    elif suffix == ".xlsx":
                        preview.update(_xlsx_preview(data))
                    elif suffix == ".mat":
                        preview.update(_mat_preview(data))
                    elif suffix == ".xls":
                        preview.update({"kind": "xls", "inspection": "legacy XLS; mapping will require xlrd"})
                else:
                    preview["inspection"] = "metadata only: member exceeds schema-preview byte limit"
                candidate_tables.append(preview)
    return {
        "member_count": len(members),
        "total_uncompressed_bytes": sum(item["uncompressed_bytes"] for item in members),
        "members": members,
        "candidate_tables": candidate_tables,
    }


def inventory(source_dir: Path, source_manifest_path: Path) -> dict:
    identity = verify_source_dir(source_dir, source_manifest_path, require_readme=True)
    bundle = source_dir / REQUIRED_BUNDLE
    schema = inspect_bundle(bundle)
    return {
        "schema_version": 1,
        "status": "SCHEMA_ONLY_INVENTORY_COMPLETE_NO_METHOD_OUTCOMES",
        "boundary": "No tracking-method performance metrics are computed by this command.",
        "source_identity": identity,
        "archive": schema,
        "next_gate": "Review this inventory and commit docs/dryad-confirmatory-schema-lock.json with status LOCKED before normalization/evaluation.",
    }


def _read_text_records(archive: zipfile.ZipFile, member: str, spec: dict) -> list[dict[str, Any]]:
    raw = archive.read(member).decode("utf-8-sig", errors="strict")
    delimiter = spec.get("delimiter")
    if not delimiter:
        suffix = Path(member).suffix.lower()
        delimiter = "\t" if suffix == ".tsv" else ","
    reader = csv.reader(io.StringIO(raw), delimiter=str(delimiter))
    rows = list(reader)
    header_row = int(spec.get("header_row", 1)) - 1
    if header_row < 0 or header_row >= len(rows):
        raise ValueError(f"invalid header_row for {member}")
    header = [str(value).strip() for value in rows[header_row]]
    records = []
    for values in rows[header_row + 1 :]:
        if not any(str(value).strip() for value in values):
            continue
        records.append({key: values[index] if index < len(values) else "" for index, key in enumerate(header)})
    return records


def _read_xlsx_records(archive: zipfile.ZipFile, member: str, spec: dict) -> list[dict[str, Any]]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise ValueError("XLSX normalization requires `pip install openpyxl`") from exc
    workbook = openpyxl.load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
    sheet_name = spec.get("sheet")
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_row = int(spec.get("header_row", 1)) - 1
    if header_row < 0 or header_row >= len(rows):
        raise ValueError(f"invalid header_row for {member}")
    header = ["" if value is None else str(value).strip() for value in rows[header_row]]
    records = []
    for values in rows[header_row + 1 :]:
        if not any(value is not None and str(value).strip() for value in values):
            continue
        records.append({key: values[index] if index < len(values) else None for index, key in enumerate(header)})
    return records


def _read_mat_records(archive: zipfile.ZipFile, member: str, spec: dict) -> list[list[Any]]:
    try:
        import scipy.io  # type: ignore
    except ImportError as exc:
        raise ValueError("MAT normalization requires `pip install scipy`") from exc
    variable = spec.get("matlab_variable")
    if not variable:
        raise ValueError(f"matlab_variable must be frozen for {member}")
    payload = scipy.io.loadmat(io.BytesIO(archive.read(member)), squeeze_me=False)
    if variable not in payload:
        raise ValueError(f"MAT variable {variable!r} not found in {member}")
    array = payload[variable]
    if getattr(array, "ndim", None) != 2:
        raise ValueError(f"MAT variable {variable!r} must be a 2D numeric matrix")
    return array.tolist()


def _value(record: Any, selector: Any, field: str) -> Any:
    if selector is None:
        return None
    if isinstance(record, dict):
        if not isinstance(selector, str) or selector not in record:
            raise ValueError(f"column selector for {field} is not present: {selector!r}")
        return record[selector]
    if not isinstance(selector, int):
        raise ValueError(f"MAT column selector for {field} must be an integer")
    if selector < 0 or selector >= len(record):
        raise ValueError(f"MAT column selector out of range for {field}: {selector}")
    return record[selector]


def _finite_float(value: Any, field: str) -> float:
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"non-finite {field}")
    return result


def _normalize_records(records: list[Any], experiment_id: str, spec: dict) -> tuple[list[dict], dict]:
    mapping = spec.get("column_map", {})
    for required in ("track_id", "x_um", "y_um"):
        if mapping.get(required) is None:
            raise ValueError(f"schema lock missing {required} mapping for {experiment_id}")
    if mapping.get("frame") is None and mapping.get("time_min") is None:
        raise ValueError(f"schema lock needs frame or time mapping for {experiment_id}")
    cell_filter = spec.get("cell_type_filter", {})
    accepted = {str(value).strip().lower() for value in cell_filter.get("accepted_values", [])}
    if cell_filter.get("enabled") and not accepted:
        raise ValueError(f"enabled cell-type filter has no accepted values for {experiment_id}")
    coord_scale = float(spec.get("coordinate_scale_to_um", 1.0))
    time_scale = float(spec.get("time_scale_to_min", 1.0))
    fallback_dt = spec.get("frame_interval_min_if_time_absent")
    if fallback_dt is not None:
        fallback_dt = float(fallback_dt)
        if fallback_dt <= 0:
            raise ValueError("frame_interval_min_if_time_absent must be positive")

    intermediate = []
    excluded_cell_type = 0
    for record in records:
        if cell_filter.get("enabled"):
            selector = mapping.get("cell_type")
            if selector is None:
                raise ValueError(f"cell-type filter enabled but cell_type mapping missing for {experiment_id}")
            source_cell_type = str(_value(record, selector, "cell_type")).strip().lower()
            if source_cell_type not in accepted:
                excluded_cell_type += 1
                continue
        track_id = str(_value(record, mapping["track_id"], "track_id")).strip()
        if not track_id:
            raise ValueError(f"empty track_id in {experiment_id}")
        frame_value = _value(record, mapping.get("frame"), "frame")
        time_value = _value(record, mapping.get("time_min"), "time_min")
        frame = int(round(float(frame_value))) if frame_value not in (None, "") else None
        time_min = _finite_float(time_value, "time_min") * time_scale if time_value not in (None, "") else None
        intermediate.append({
            "track_id": track_id,
            "frame": frame,
            "time_min": time_min,
            "x_um": _finite_float(_value(record, mapping["x_um"], "x_um"), "x_um") * coord_scale,
            "y_um": _finite_float(_value(record, mapping["y_um"], "y_um"), "y_um") * coord_scale,
        })

    if not intermediate:
        raise ValueError(f"no glioma observations retained for {experiment_id}")
    if any(row["frame"] is None for row in intermediate):
        if fallback_dt is None:
            raise ValueError(f"cannot derive frame without frozen frame interval for {experiment_id}")
        min_time = min(float(row["time_min"]) for row in intermediate)
        for row in intermediate:
            ratio = (float(row["time_min"]) - min_time) / fallback_dt
            frame = int(round(ratio))
            if abs(ratio - frame) > 1e-5:
                raise ValueError(f"time values do not map cleanly to frames in {experiment_id}")
            row["frame"] = frame
    if any(row["time_min"] is None for row in intermediate):
        if fallback_dt is None:
            raise ValueError(f"cannot derive time without frozen frame interval for {experiment_id}")
        for row in intermediate:
            row["time_min"] = int(row["frame"]) * fallback_dt

    seen = set()
    by_track: dict[str, list[dict]] = {}
    for row in intermediate:
        frame = int(row["frame"])
        key = (row["track_id"], frame)
        if key in seen:
            raise ValueError(f"duplicate track/frame {key} in {experiment_id}")
        seen.add(key)
        row["frame"] = frame
        by_track.setdefault(row["track_id"], []).append(row)
    for track_id, track_rows in by_track.items():
        track_rows.sort(key=lambda item: item["frame"])
        for left, right in zip(track_rows, track_rows[1:]):
            if float(right["time_min"]) <= float(left["time_min"]):
                raise ValueError(f"non-increasing time in track {track_id} of {experiment_id}")

    rows = [
        {
            "experiment_id": experiment_id,
            "cell_type": "glioma",
            "track_id": row["track_id"],
            "frame": row["frame"],
            "time_min": float(row["time_min"]),
            "x_um": float(row["x_um"]),
            "y_um": float(row["y_um"]),
        }
        for row in intermediate
    ]
    rows.sort(key=lambda item: (item["frame"], item["track_id"]))
    return rows, {
        "source_records": len(records),
        "retained_glioma_observations": len(rows),
        "excluded_by_cell_type_filter": excluded_cell_type,
        "track_count": len(by_track),
        "frame_min": min(row["frame"] for row in rows),
        "frame_max": max(row["frame"] for row in rows),
    }


def _records_for_member(archive: zipfile.ZipFile, member: str, spec: dict) -> list[Any]:
    names = set(archive.namelist())
    if member not in names:
        raise ValueError(f"frozen source member not found: {member}")
    fmt = (spec.get("format") or Path(member).suffix.lower().lstrip(".")).lower()
    if fmt in {"csv", "tsv", "txt", "dat"}:
        return _read_text_records(archive, member, spec)
    if fmt == "xlsx":
        return _read_xlsx_records(archive, member, spec)
    if fmt == "mat":
        return _read_mat_records(archive, member, spec)
    raise ValueError(f"unsupported frozen source format {fmt!r} for {member}")


def normalize(source_dir: Path, source_manifest_path: Path, schema_lock_path: Path, output_dir: Path) -> dict:
    lock = _load_json(schema_lock_path)
    if lock.get("status") != "LOCKED":
        raise ValueError("Dryad schema lock is not LOCKED; evaluation boundary preserved")
    identity = verify_source_dir(source_dir, source_manifest_path, require_readme=True)
    if lock.get("source_bundle", {}).get("md5") != REQUIRED_BUNDLE_MD5:
        raise ValueError("schema lock source identity drifted")
    experiments = lock.get("experiments", {})
    if sorted(experiments) != ["experiment_1", "experiment_2", "experiment_3"]:
        raise ValueError("schema lock must define exactly experiment_1, experiment_2 and experiment_3")

    output_dir.mkdir(parents=True, exist_ok=True)
    normalized = {}
    with zipfile.ZipFile(source_dir / REQUIRED_BUNDLE) as archive:
        for info in archive.infolist():
            _validate_member(info)
        for experiment_id in sorted(experiments):
            spec = experiments[experiment_id]
            source_members = spec.get("source_members", [])
            if not source_members:
                raise ValueError(f"no frozen source_members for {experiment_id}")
            records: list[Any] = []
            for member in source_members:
                records.extend(_records_for_member(archive, str(member), spec))
            rows, summary = _normalize_records(records, experiment_id, spec)
            output_path = output_dir / f"{experiment_id}.csv"
            with output_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=CANONICAL_COLUMNS)
                writer.writeheader()
                writer.writerows(rows)
            normalized[experiment_id] = {
                **summary,
                "path": output_path.name,
                "sha256": _digest(output_path, "sha256"),
                "source_members": source_members,
            }

    manifest = {
        "schema_version": 1,
        "status": "NORMALIZED_THREE_EXPERIMENT_DRYAD_REFERENCE",
        "source_identity": identity,
        "schema_lock_sha256": _digest(schema_lock_path, "sha256"),
        "python": platform.python_version(),
        "experiments": normalized,
        "biological_n": 3,
        "boundary": "Normalization only; no tracking-method performance metrics computed.",
    }
    manifest_path = output_dir / "normalized-manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return manifest


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    inventory_parser = subparsers.add_parser("inventory", help="Verify deposited hashes and create schema-only inventory")
    inventory_parser.add_argument("--source-dir", type=Path, required=True)
    inventory_parser.add_argument("--source-manifest", type=Path, default=Path("docs/external-biological-context-source-manifest.json"))
    inventory_parser.add_argument("--output", type=Path, required=True)

    normalize_parser = subparsers.add_parser("normalize", help="Normalize all three experiments after the schema lock is committed")
    normalize_parser.add_argument("--source-dir", type=Path, required=True)
    normalize_parser.add_argument("--source-manifest", type=Path, default=Path("docs/external-biological-context-source-manifest.json"))
    normalize_parser.add_argument("--schema-lock", type=Path, default=Path("docs/dryad-confirmatory-schema-lock.json"))
    normalize_parser.add_argument("--output-dir", type=Path, required=True)

    args = parser.parse_args(argv)
    try:
        if args.command == "inventory":
            result = inventory(args.source_dir, args.source_manifest)
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
            print(f"Wrote schema-only inventory: {args.output}")
        else:
            result = normalize(args.source_dir, args.source_manifest, args.schema_lock, args.output_dir)
            print(json.dumps(result, indent=2, ensure_ascii=False))
    except (OSError, ValueError, json.JSONDecodeError, zipfile.BadZipFile) as exc:
        print(f"Dryad local pipeline failed: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
